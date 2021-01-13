import requests
import settings
import aux_functions
import regexLH
import pytesseract
import random
import os
from os import walk
from PIL import Image
import shutil
from tika import parser
import json
import sys
from docx import Document
import xlrd
import hashlib
import mimetypes
try:
    from xml.etree.cElementTree import XML
except ImportError:
    from xml.etree.ElementTree import XML
import zipfile

from ibm_watson import VisualRecognitionV3
from ibm_cloud_sdk_core.authenticators import IAMAuthenticator

iam = IAMAuthenticator(settings.WATSON_API_KEY)
visual_recognition = VisualRecognitionV3(
    version=settings.WATSON_VERSION,
    authenticator=iam
)

visual_recognition.set_service_url(settings.WATSON_SERVICE_URL)

WORD_NAMESPACE = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
PARA = WORD_NAMESPACE + 'p'
TEXT = WORD_NAMESPACE + 't'

try:
    from PIL import Image
except ImportError:
    import Image

# def readWord(filename):
#     try:
#         document = Document(filename)
#         print(document)
#     except Exception as e:
#         print(e)

def readImage(fileOutput):
    try:
        file = fileOutput['name']
        url = fileOutput['url']
        # print(url)
        if not file.endswith('.png') and not file.endswith('.jpg') and not file.endswith('.jpeg') and not file.endswith('.gif') and not file.endswith('.tif'):
            return;
        loadAs = os.path.join(aux_functions.TmpFolder, fileOutput['name'])
        with open(loadAs,'rb') as images_file:
            classes_result = visual_recognition.classify(
            images_file=images_file,
            threshold=settings.WATSONTHRESHOLD,
            classifier_ids=[settings.WATSONCLASSIFIERID]
            ).get_result()
        classes = []
        classes = classes_result['images'][0]['classifiers'][0]['classes']
        for classifier in classes:
            image = {
                'url':url,
                'type':classifier['class'],
                'score':classifier['score']
            }
            return image;
    except Exception as e:
        print('readImage:',e)
        return;


def readWord(filename):
    try:
        document = zipfile.ZipFile(filename)
        xml_content = document.read('word/document.xml')
        document.close()
        tree = XML(xml_content)

        paragraphs = '';
        for paragraph in tree.getiterator(PARA):
            texts = [node.text
                     for node in paragraph.getiterator(TEXT)
                     if node.text]
            if texts:
                paragraphs += str(texts) + ',';
        return paragraphs
    except Exception as e:
        print('ReadWord exception',e)

def readExcel(filename):
    try:
        loc = (filename)
        # To open Workbook
        wb = xlrd.open_workbook(loc)
        sheet = wb.sheet_by_index(0)
        # read header values into the list
        #keys = [sheet.cell(0, col_index).value for col_index in range(sheet.ncols)]
        row_list = '';
        for row_index in range(1, sheet.nrows):
            for col_index in range(sheet.ncols):
                d = str(sheet.cell(row_index, col_index).value)
                row_list += d + '<>';
        #print(row_list)
        return row_list;
    except Exception as e:
        print('ReadExcel exception',e)

def ocr_core(filename): #extract text from images
    try:
        text = pytesseract.image_to_string(Image.open(filename))  # We'll use Pillow's Image class to open the image and pytesseract to detect the string in the image
        return text;
    except Exception as e:
        print('OCR exception',e)

def readPDF(filename):
    try:
        raw = parser.from_file(filename)
        #print(raw['content'])
        return raw['content'];
    except Exception as e:
        print('Readpdf exception:',e)

def download(url,result):
    headers = {'user-agent': settings.USERAGENT}
    r = requests.get(url,headers=headers)
    content_type = r.headers['content-type']
    extension = mimetypes.guess_extension(content_type,True)
    extension2 = mimetypes.guess_type(url,True)
    print(extension,extension2,url)
    if extension == '.jpe':
        extension = '.jpg';
    if extension:
        filename = str(random.randint(1,99999)) + extension;
        saveAs = os.path.join(aux_functions.TmpFolder, filename)
        # print(filename)
        with open(saveAs,'wb') as f:
                #write the contents of the response (r.content)
                # to a new file in binary mode.
                f.write(r.content)
        sha1 = aux_functions.getHash(saveAs)
        if result['fileSHAs'].count(sha1) == 0:
            file = {
                'name':filename,
                'scanned':False,
                'url':url,
                'sha1':sha1,
                'type':extension,
                'text':''
            }
            return file;
        else:
            aux_functions.deleteOneFile(saveAs)

fileTypePhoto = ['.png','.jpg','.gif','.tif','.jpeg']

def validImage(file):
    loadAs = os.path.join(aux_functions.TmpFolder, file['name'])
    im = Image.open(loadAs)
    print(im.size)
    if im.size[0] > 80 and im.size[1] > 80:
        return True
    else:
        return False

def readAllFiles(result):
    allOutput = [];
    countFiles = 0;
    index = 0;
    while index < len(result['files']): #while index loop not necessary needed TODO work on simplifying this
        file = result['files'][index]['file']
        file['scanned'] = True;
        if file['type'] in fileTypePhoto and validImage(file):
            imageClass = readImage(file)
            result['cost'] += 0.002
            if not imageClass == None:
                file['watsonDetection'] = imageClass
                result['documents'][imageClass['type']].append(imageClass)
            file['text'] = ocr_core(aux_functions.TmpFolder+file['name'])
        elif file['type'] == '.pdf':
            file['text'] = readPDF(aux_functions.TmpFolder+file['name'])
        elif file['type'] == '.xlsx':
            file['text'] = readExcel(aux_functions.TmpFolder+file['name'])
        elif file['type'] == '.docx':
            file['text'] = readWord(aux_functions.TmpFolder+file['name'])
        elif file['type'] == '.doc':
            file['text'] = readWord(aux_functions.TmpFolder+file['name'])
        else:
            file['scanned'] = False;
        index += 1;
    #print(allOutput[1].url)
    return result;
